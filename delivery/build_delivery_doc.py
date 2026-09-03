from __future__ import annotations

from pathlib import Path

from build_diagrams import build_diagrams
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "项目说明书.md"
OUTPUT = ROOT / "AgentBI项目说明文档-评审版.docx"
BLUE = "2563EB"
NAVY = "0B2545"
MUTED = "64748B"
LIGHT = "E8EEF5"


def set_font(run, name="Microsoft YaHei", size=11, bold=False, color="111827"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("InsightPilot AgentBI | AI Coding 大赛交付文档")
    set_font(r, size=8.5, color=MUTED)


build_diagrams()
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)
add_footer(section)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 13, BLUE, 11, 5),
    ("Heading 3", 11.5, NAVY, 8, 4),
]:
    style = styles[name]
    style.font.name = "Microsoft YaHei"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(42)
p.paragraph_format.space_after = Pt(8)
r = p.add_run("INSIGHTPILOT")
set_font(r, size=11, bold=True, color=BLUE)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("AgentBI 项目说明书")
set_font(r, size=28, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(22)
r = p.add_run("可验证的经营分析 Agent · 真实查询 · 可视化下钻 · 审计报告")
set_font(r, size=13, color=MUTED)

table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = False
labels = [("项目名称", "InsightPilot AgentBI"), ("交付类型", "AI Coding 大赛参赛作品"), ("文档版本", "1.0"), ("生成日期", "2026-09-03（北京时间）")]
for row, (label, value) in zip(table.rows, labels):
    set_cell_width(row.cells[0], 2400)
    set_cell_width(row.cells[1], 6960)
    shade(row.cells[0], LIGHT)
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_font(row.cells[0].paragraphs[0].add_run(label), size=9.5, bold=True, color=NAVY)
    set_font(row.cells[1].paragraphs[0].add_run(value), size=9.5)

doc.add_section(WD_SECTION.NEW_PAGE)
for raw in SOURCE.read_text(encoding="utf-8").splitlines():
    line = raw.strip()
    if not line or line == "# InsightPilot AgentBI 项目说明书":
        continue
    if line.startswith("### "):
        doc.add_paragraph(line[4:], style="Heading 3")
    elif line.startswith("## "):
        doc.add_paragraph(line[3:], style="Heading 1")
    elif line.startswith("![") and "](" in line and line.endswith(")"):
        alt = line[2 : line.index("](")]
        relative = line[line.index("](") + 2 : -1]
        picture = doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_before = Pt(6)
        picture.paragraph_format.space_after = Pt(3)
        picture.add_run().add_picture(str(ROOT / relative), width=Inches(6.45))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        set_font(caption.add_run(alt), size=9, color=MUTED)
    elif line[:3].rstrip(".").isdigit() and ". " in line:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        r = p.add_run(line.split(". ", 1)[1])
        set_font(r, size=10.5)
    else:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0.28)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(line)
        set_font(r, size=10.5)

doc.core_properties.title = "InsightPilot AgentBI 项目说明文档（评审版）"
doc.core_properties.subject = "AI Coding 大赛交付文档"
doc.core_properties.author = "InsightPilot AgentBI 项目组"
doc.save(OUTPUT)
print(OUTPUT)
